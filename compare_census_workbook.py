#!/usr/bin/env python3
"""Compare a manually-researched census workbook against the live V26 portal.

Additive, read-only comparison: loads the "All Canada" sheet of a census
workbook (columns: Funeral Home / Location, City, Province, Street Address,
Postal Code, Phone, Email, Website, Contact Person / Manager, Type / Notes,
Source URL, Verified / Source Date) and matches each row against
`data/portal_findings.json` by normalized website domain first, then by
normalized (name, city, province). Rows that match neither are written out as
new-business candidates for further review/crawling. Performs no crawling,
no DNS resolution, and no network access.

Match confidence is intentionally two-tiered: a domain match is strong
evidence of the same physical business regardless of name spelling; a
name+city+province match catches the common case where a workbook row has no
usable website. Everything else becomes a candidate, sorted by whether a
website is present (crawlable) or not (phone/regulator-only).
"""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from pathlib import Path
from typing import Optional

import openpyxl

from discovery.ingestion import domain_from_website
from discovery.province_mapping import normalize_province

SOURCE_XLSX = Path("data/discovery_sources/census_2026_09_10_v52.xlsx")
SOURCE_SHEET = "All Canada"
PORTAL_SOURCE = Path("data/portal_findings.json")
OUTPUT = Path("data/generated/enrichment/census_2026_09_comparison.json")
EXPECTED_PORTAL_VERSION = "V26"

CENSUS_COLUMNS = (
    "name", "city", "province_raw", "street_address", "postal_code",
    "phone", "email", "website", "contact", "notes", "source_url",
    "verified_date",
)

NAME_NOISE_WORDS = re.compile(
    r"\b(FUNERAL HOME|FUNERAL HOMES|FUNERAL SERVICES?|FUNERAL CHAPEL|"
    r"FUNERAL SERVICE|CREMATORIUM|CREMATION|LTD|INC|CO|CORP|LIMITED|THE|AND)\b"
)


class SourceValidationError(ValueError):
    """The workbook or portal source does not meet the required shape."""


def normalize_name(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9 &]", " ", name or "").upper()
    cleaned = NAME_NOISE_WORDS.sub(" ", cleaned)
    cleaned = cleaned.replace("&", " ")
    return " ".join(cleaned.split())


def load_census_rows(xlsx_path: Path, sheet_name: str = SOURCE_SHEET) -> list[dict]:
    if not xlsx_path.is_file():
        raise SourceValidationError(f"Census workbook not found: {xlsx_path}")
    try:
        workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises various concrete errors
        raise SourceValidationError(f"Could not open workbook: {exc}") from exc
    try:
        if sheet_name not in workbook.sheetnames:
            raise SourceValidationError(
                f"Workbook has no {sheet_name!r} sheet (found {workbook.sheetnames!r})"
            )
        sheet = workbook[sheet_name]

        rows: list[dict] = []
        row_iter = sheet.iter_rows(values_only=True)
        header = next(row_iter, None)
        if header is None:
            raise SourceValidationError(f"Sheet {sheet_name!r} has no header row")

        for raw in row_iter:
            if raw is None or all(cell is None for cell in raw):
                continue
            values = list(raw) + [None] * (len(CENSUS_COLUMNS) - len(raw))
            row = dict(zip(CENSUS_COLUMNS, values))
            row = {key: (str(value).strip() if value is not None else "") for key, value in row.items()}
            if not row["name"]:
                continue
            row["province"] = normalize_province(row["province_raw"]) or ""
            row["domain"] = domain_from_website(row["website"]) if row["website"] else ""
            rows.append(row)
        return rows
    finally:
        workbook.close()


def load_portal_records(path: Path) -> tuple[list[dict], str]:
    if not path.is_file():
        raise SourceValidationError(f"Portal source not found: {path}")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise SourceValidationError(f"Portal source is not valid JSON: {exc}") from exc
    if not isinstance(payload, dict):
        raise SourceValidationError("Portal source must be a JSON object")
    if payload.get("version") != EXPECTED_PORTAL_VERSION:
        raise SourceValidationError(
            f"Portal source version is {payload.get('version')!r}, "
            f"expected {EXPECTED_PORTAL_VERSION!r}"
        )
    records = payload.get("records")
    if not isinstance(records, list):
        raise SourceValidationError("Portal source 'records' must be a list")
    return records, payload["version"]


def build_portal_indexes(records: list[dict]) -> tuple[dict, dict]:
    by_domain: dict[str, list[str]] = {}
    by_namecity: dict[tuple, list[str]] = {}
    for record in records:
        record_id = record.get("directory_record_id", "")
        website = record.get("website") or ""
        domain = domain_from_website(website) if website else ""
        if domain:
            by_domain.setdefault(domain, []).append(record_id)
        key = (
            normalize_name(record.get("company", "")),
            (record.get("city") or "").strip().upper(),
            record.get("province", ""),
        )
        by_namecity.setdefault(key, []).append(record_id)
    return by_domain, by_namecity


def compare(census_rows: list[dict], portal_records: list[dict]) -> dict:
    by_domain, by_namecity = build_portal_indexes(portal_records)

    domain_matches = []
    namecity_matches = []
    new_candidates = []

    for row in census_rows:
        if row["domain"] and row["domain"] in by_domain:
            domain_matches.append({"census_row": row, "portal_record_ids": by_domain[row["domain"]]})
            continue
        key = (normalize_name(row["name"]), row["city"].strip().upper(), row["province"])
        if key in by_namecity:
            namecity_matches.append({"census_row": row, "portal_record_ids": by_namecity[key]})
            continue
        new_candidates.append(row)

    new_candidates.sort(key=lambda row: (row["province"], not row["website"], row["name"]))

    summary = {
        "census_rows": len(census_rows),
        "portal_records": len(portal_records),
        "domain_matches": len(domain_matches),
        "namecity_matches": len(namecity_matches),
        "new_candidates": len(new_candidates),
        "new_candidates_with_website": sum(1 for row in new_candidates if row["website"]),
        "new_candidates_without_website": sum(1 for row in new_candidates if not row["website"]),
        "new_candidates_by_province": dict(
            sorted(Counter(row["province"] for row in new_candidates).items())
        ),
    }
    return {
        "summary": summary,
        "domain_matches": domain_matches,
        "namecity_matches": namecity_matches,
        "new_candidates": new_candidates,
    }


def write_comparison(
    xlsx_path: Path, portal_path: Path, output: Path, sheet_name: str = SOURCE_SHEET
) -> dict:
    census_rows = load_census_rows(xlsx_path, sheet_name)
    portal_records, _version = load_portal_records(portal_path)
    result = compare(census_rows, portal_records)

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    return result["summary"]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, default=SOURCE_XLSX)
    parser.add_argument("--sheet", default=SOURCE_SHEET)
    parser.add_argument("--portal", type=Path, default=PORTAL_SOURCE)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()

    try:
        summary = write_comparison(args.xlsx, args.portal, args.output, args.sheet)
    except SourceValidationError as exc:
        print(f"REFUSED: {exc}")
        return 1

    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
